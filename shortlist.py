"""Build the daily list of 20 internship applications.

Each run hands out PER_DAY roles that were never handed out before, one per
company, from every open posting in the four areas (quant, data science / ML,
AI consulting, internet backend). Order (see PRIORITIES; "preferred dates"
means the posting states dates from June 1 on):
  1. Manhattan / Jersey City in person, preferred dates
  2. remote, preferred dates
  3. Manhattan / Jersey City, other or unstated dates
  4. other city in person, preferred dates
  5. everything else (remote, then other cities)
Inside each: the four areas take turns, stated dates before unstated, newest first.

Dropped before the pick (the applicant is a green-card holder, undergraduate,
graduating May 2028, and wants SUMMER internships only):
  - roles that need US citizenship or a security clearance
  - PhD / Master's / MBA-only roles: by title, and by the posting text
    ("currently pursuing a Master's or PhD")
  - roles whose posting states a graduation window that May 2028 misses
  - postings the job board has taken down (the tracker can lag a day or two)
  - terms that have already started (for example Fall 2026)
  - roles not confirmed as a summer internship (title, season label, or the
    posting text must say so)
  - roles whose stated dates fall outside WINDOW_START .. WINDOW_END
    (May 12 - Aug 31, a hard cut: a stated start of May 11 or earlier is
    out, and so is a start given only as "May"). Most postings state no dates
    and are kept as plain "Summer 2027".
Posting text checks come from data/term_cache.json, written by term_dates.py.

Applied log: every role that lands on a daily list is written to the applied
log that day and is never handed out again. On each run, logged roles whose
posting has closed (expired, or no longer an internship) are removed from it.
Roles you did not apply to can be taken out with --skip so they come back.

Reads  docs/api/jobs.json           (written by `run.py update`)
Writes shortlists/<date>.xlsx       (the 20; <date> like 2026-9-8, local date)
       shortlists/applied log.xlsx  (the applied log, readable copy)
       data/today20.json            (same 20, read by the open-in-Chrome command)
       data/shortlist_latest.json   (same 20, for send_email.py --test)
       data/shortlist_new.json      (same 20, queued for send_email.py)
Keeps  data/applied_log.json        (the applied log itself)

Run:   .venv\\Scripts\\python.exe shortlist.py            today's 20 (same list if run again today)
       .venv\\Scripts\\python.exe shortlist.py --new      another 20 today (written as "<date> batch 2.xlsx")
       .venv\\Scripts\\python.exe shortlist.py --dry-run  show the pick, change nothing
       .venv\\Scripts\\python.exe shortlist.py --urls     print the links of today's 20
       .venv\\Scripts\\python.exe shortlist.py --skip 3 7  take #3 and #7 of today's list out of the log
"""

import json
import os
import re
import sys
from collections import Counter, OrderedDict
from datetime import date, datetime, timedelta

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

PER_DAY = 20         # roles handed out per list
WINDOW_START = (5, 12)   # internship may not start before May 12
WINDOW_END = (8, 31)     # and may not end after August 31
PREFERRED_START = (6, 1)  # stated dates from June 1 on fill the list first
KEEP_FILES_DAYS = 5  # dated workbooks older than this are deleted

ROOT = os.path.dirname(os.path.abspath(__file__))
OUT_DIR = os.path.join(ROOT, "shortlists")
JOBS = os.path.join(ROOT, "docs", "api", "jobs.json")
LOG = os.path.join(ROOT, "data", "applied_log.json")
LOG_XLSX = os.path.join(OUT_DIR, "applied log.xlsx")
TODAY = os.path.join(ROOT, "data", "today20.json")
NEW = os.path.join(ROOT, "data", "shortlist_new.json")
LATEST = os.path.join(ROOT, "data", "shortlist_latest.json")
TERM_CACHE = os.path.join(ROOT, "data", "term_cache.json")

GROUPS = OrderedDict([
    ("quant", "Quant"),
    ("data_science", "Data Science / ML / AI"),
    ("ai_consulting", "AI & Data Consulting"),
    ("backend", "Internet / Backend Software"),
])
TIERS = OrderedDict([
    ("nyc", "Manhattan / Jersey City"),
    ("remote", "Remote / online"),
    ("other", "Other US cities"),
])

CONSULTING_FIRMS = [
    "booz allen", "accenture", "deloitte", "mckinsey", "quantumblack", "bcg", "boston consulting",
    "ernst", " ey ", "pwc", "pricewaterhouse", "kpmg", "bain", "capgemini", "ibm", "slalom",
    "zs associates", "west monroe", "oliver wyman", "guidehouse", "cognizant", "infosys", "wipro",
    "charles river associates", "publicis sapient", "thoughtworks", "epam", "globant",
    "analysis group", "cornerstone research", "nera", "brattle", "fti consulting", "alvarez",
    "mitre", "saic", "leidos", "caci", "general dynamics information technology",
]
INTERNET_COS = [
    "meta", "google", "alphabet", "youtube", "amazon", "netflix", "uber", "airbnb", "stripe",
    "shopify", "pinterest", "snap", "reddit", "doordash", "instacart", "lyft", "dropbox", "twitch",
    "roblox", "cloudflare", "datadog", "databricks", "snowflake", "discord", "coinbase",
    "robinhood", "block", "paypal", "ebay", "etsy", "zillow", "yelp", "linkedin", "microsoft",
    "salesforce", "adobe", "atlassian", "twilio", "okta", "mongodb", "hashicorp", "vercel",
    "notion", "figma", "canva", "duolingo", "spotify", "zoom", "slack", "expedia", "booking",
    "tiktok", "bytedance", "tinder", "match group", "chime", "affirm", "plaid", "brex", "ramp",
    "rippling", "gusto", "wayfair", "chewy", "grubhub", "nextdoor", "quora", "squarespace", "wix",
    "godaddy", "akamai", "fastly", "digitalocean", "confluent", "elastic", "grafana", "hubspot",
    "servicenow", "workday", "intuit", "docusign", "box", "asana", "airtable", "webflow",
    "anthropic", "openai", "perplexity", "scale ai", "cohere", "mistral", "hugging face",
    "whatnot", "faire", "flexport", "samsara", "toast", "carta", "mercury", "sofi", "kraken",
    "gemini", "ripple", "opensea", "alchemy",
]

RE_QUANT = re.compile(r"\bquant|trading|trader|market[- ]mak|systematic|alpha research|portfolio", re.I)
RE_DS = re.compile(
    r"data scien|data analy|analytics|machine learning|\bML\b|\bAI\b|artificial intelligence|"
    r"deep learning|\bNLP\b|computer vision|statistic|applied scien|research scien|\bLLM|"
    r"generative|data engineer|business intelligence|\bBI\b|decision scien|forecast",
    re.I,
)
RE_BACKEND_ANY = re.compile(
    r"back[- ]?end|server[- ]side|platform|infrastructure|distributed|\bAPI\b|cloud|"
    r"full[- ]?stack|site reliability|\bSRE\b|devops|microservice|\bweb\b|internet|"
    r"scalab|systems engineer|software engineer|software develop|\bSWE\b|\bSDE\b",
    re.I,
)
RE_BACKEND_EXPLICIT = re.compile(
    r"back[- ]?end|server|platform|infrastructure|distributed|\bAPI\b|cloud|"
    r"full[- ]?stack|site reliability|\bSRE\b|devops|microservice|\bweb\b", re.I,
)
RE_EXCLUDE_BACKEND = re.compile(
    r"front[- ]?end|mobile|ios\b|android|embedded|firmware|hardware|fpga|asic|rtl|"
    r"qa\b|test engineer|game|roadway|water|civil|construction|transportation|structural|"
    r"service desk|project support", re.I,
)
# "NY" means Manhattan; Jersey City is fine too. Brooklyn, Hoboken etc. are not.
RE_NYC = re.compile(r"\bnyc\b|manhattan|jersey city|new york city", re.I)
RE_LOC_PREFIX = re.compile(r"^(?:office|hybrid|remote|onsite|on-site|us|usa)\s*[-:]\s*", re.I)
RE_NY_STATE = re.compile(r"^(?:ny\b|new york)", re.I)
BIG_CITIES = {"chicago", "san francisco", "seattle", "boston", "austin", "los angeles", "london",
              "houston", "miami", "philadelphia", "washington", "atlanta", "dallas", "denver"}
RE_REMOTE = re.compile(r"remote|all locations|virtual|work from home|anywhere", re.I)
RE_DATED = re.compile(r"^(\d{4})-(\d{1,2})-(\d{1,2})[ .]")

# Eligibility, judged from the title (the posting text is judged by the tracker
# itself and arrives as the `sponsorship` field).
RE_CITIZEN_TITLE = re.compile(r"clearance|ts/?sci|top secret|u\.?s\.? citizen", re.I)
RE_GRAD_ONLY = re.compile(
    r"ph\.?\s?d\b|doctora|\bmba\b|master|\bgraduate\b|\bm\.?s\.?\b|postdoc|\bjd\b|law student", re.I)
RE_UNDERGRAD_OK = re.compile(r"bachelor|undergrad|\bb\.?s\.?\b|\bb\.?a\.?\b", re.I)
RE_TERM = re.compile(r"\b(spring|winter|summer|fall|autumn)\s+(20\d\d)\b", re.I)
TERM_START_MONTH = {"winter": 1, "spring": 1, "summer": 5, "fall": 8, "autumn": 8}

NOTES = {"us-persons": "citizens or green card OK"}


# ----------------------------------------------------------------- classify

def is_consulting(j):
    c = " " + j["company"].lower() + " "
    return any(f in c for f in CONSULTING_FIRMS) or "consult" in j["title"].lower()


def is_internet(j):
    c = j["company"].lower()
    return any(re.search(r"\b" + re.escape(n) + r"\b", c) for n in INTERNET_COS)


def classify(j):
    t = j["title"]
    if j.get("category") == "Quant" or RE_QUANT.search(t):
        return "quant"
    if is_consulting(j) and (RE_DS.search(t) or j.get("category") == "Data & ML/AI"):
        return "ai_consulting"
    if j.get("category") == "Data & ML/AI" or RE_DS.search(t):
        return "data_science"
    if j.get("category") == "Software" and RE_BACKEND_ANY.search(t) and not RE_EXCLUDE_BACKEND.search(t):
        if is_internet(j) or RE_BACKEND_EXPLICIT.search(t):
            return "backend"
    return None


def term_started(season, today):
    """True when 'Fall 2026'-style season has already begun."""
    m = RE_TERM.search(season or "")
    if not m:
        return False
    return date(int(m[2]), TERM_START_MONTH[m[1].lower()], 1) <= today


def why_ineligible(j, today):
    """Reason this posting is dropped before the pick, or None when it is fine."""
    t = j["title"]
    if j.get("sponsorship") == "citizens-only" or RE_CITIZEN_TITLE.search(t):
        return "needs US citizenship or clearance"
    if RE_GRAD_ONLY.search(t) and not RE_UNDERGRAD_OK.search(t):
        return "PhD / Master's / MBA only"
    seasons = j.get("seasons") or [j.get("season") or ""]
    if all(term_started(s, today) for s in seasons) and any(RE_TERM.search(s or "") for s in seasons):
        return "term already started"
    return None


def is_nyc(location):
    """Manhattan or Jersey City. "New York" counts when it is the city, not the
    state: "New York, NY" yes, "Clifton Park, New York" no."""
    for part in re.split(r"[;|•]", location or ""):
        if RE_NYC.search(part):
            return True
        tokens = [RE_LOC_PREFIX.sub("", t.strip()).lower() for t in part.split(",")]
        for i, t in enumerate(tokens):
            if not t.startswith("new york"):
                continue
            after = tokens[i + 1] if i + 1 < len(tokens) else ""
            before = tokens[i - 1] if i else ""
            if i == 0 or RE_NY_STATE.match(after) or before in BIG_CITIES:
                return True
    return False


def tier_of(rec):
    if any(is_nyc(l) for l in rec["locations"]):
        return "nyc"
    text = rec["title"] + " " + " ".join(rec["locations"])
    if "hybrid" in text.lower():
        return "other"  # hybrid means you still go to that city
    if rec.get("remote") or RE_REMOTE.search(text):
        return "remote"
    return "other"


def posted_date(j):
    return (j.get("posted_at") or j.get("first_seen_at") or "")[:10]


def role_key(company, title):
    return re.sub(r"\s+", " ", f"{company}|{title}".lower()).strip()


def collapse(rows):
    """Merge identical company+title rows, joining their locations.

    The application link kept is the New York one when the role has several
    locations, otherwise the first.
    """
    out = OrderedDict()
    for j in rows:
        key = role_key(j["company"], j["title"])
        if key not in out:
            out[key] = dict(j, locations=[], ids=[], key=key)
        rec = out[key]
        loc = j.get("location") or ""
        if loc and loc not in rec["locations"]:
            rec["locations"].append(loc)
        if is_nyc(loc) and not is_nyc(rec.get("url_location")):
            rec["url"], rec["url_location"] = j["url"], loc
        rec["ids"].append(j["id"])
        rec["remote"] = rec.get("remote") or j.get("remote")
        if posted_date(j) > posted_date(rec):
            rec["posted_at"] = j.get("posted_at") or j.get("first_seen_at")
    return list(out.values())


# ----------------------------------------------------------------- summer + dates

RE_EARLY_START = re.compile(r"(?:early|beginning\s+of)[\s-]+may", re.I)
RE_LATE_START = re.compile(r"(?:late|end\s+of)[\s-]+may", re.I)


def term_verdict(role, cache):
    """(ok, reason or dates text, rank) for one collapsed role.

    rank 0 = start and end stated, inside the preferred June 1 - Aug 31
    rank 1 = start stated, June 1 or later, no end stated
    rank 2 = dates stated, starting May 12 - 31
    rank 3 = dates not stated

    The window is a hard cut: a stated date outside it drops the role, and so
    does a start given only as "May", which cannot be shown to be May 12 or later.
    """
    entries = [cache[i] for i in role["ids"] if cache.get(i, {}).get("has_text")]
    dated = next((e for e in entries if e.get("start")), None)

    if dated:
        (sm, sd), end, quote = dated["start"], dated.get("end"), dated["quote"]
        if sm == WINDOW_START[0] and sd is None:  # "May" with no day
            if RE_LATE_START.search(quote):
                sd = 25
            elif RE_EARLY_START.search(quote):
                sd = 1
            else:
                return False, f"start given only as May, cutoff May 12 not confirmed ({quote})", None
        too_early = (sm, sd or 31) < WINDOW_START  # May 11 or earlier is out
        too_late = sm > WINDOW_END[0] or (end and end[0] > WINDOW_END[0])
        if too_early or too_late:
            return False, f"dates outside May 12 - Aug 31 ({quote})", None

    title = role["title"].lower()
    seasons = [role.get("season") or ""] + list(role.get("seasons") or [])
    is_summer = (
        any("summer" in x.lower() for x in seasons)
        or "summer" in title
        or dated is not None  # stated dates already checked to sit inside the window
        or any(e.get("summer") and not e.get("other_terms") for e in entries)
    )
    if not is_summer:
        return False, "not confirmed as a summer internship", None
    if not dated:
        return True, "", 3
    if (sm, sd or 1) >= PREFERRED_START:
        return True, quote, 0 if end else 1
    return True, quote, 2


def requirement_verdict(role, cache):
    """(True, notes) or (False, reason), from what term_dates.py read in the posting.

    Dropped: the posting asks for PhD / Master's students only, or states a
    graduation window that May 2028 falls outside. A role whose posting could
    not be read is kept with a note to check it by hand.
    """
    entries = [cache[i] for i in role["ids"] if cache.get(i, {}).get("has_text")]
    if not entries:
        return True, ["requirements not checked (posting unreadable)"]
    degrees = {e.get("degree") for e in entries}
    if "grad-only" in degrees and "undergrad-ok" not in degrees:
        return False, "PhD / Master's only (says so in the posting)"
    grads = [e.get("grad_ok") for e in entries if e.get("grad_ok") is not None]
    if grads and not any(grads):
        return False, "graduation window in the posting misses May 2028"
    notes = []
    if any(e.get("grad_preferred") for e in entries):
        notes.append("PhD / Master's preferred, undergrads considered")
    return True, notes


# ----------------------------------------------------------------- pick

PRIORITIES = [  # the order roles are handed out in
    "1. Manhattan / Jersey City in person, preferred dates",
    "2. Remote, preferred dates",
    "3. Manhattan / Jersey City, other or unstated dates",
    "4. Other city in person, preferred dates",
    "5. Remote, other or unstated dates",
    "5. Other city, other or unstated dates",
]


def priority(r):
    """Index into PRIORITIES. Preferred dates = stated dates from June 1 on (rank 0 or 1)."""
    preferred = r.get("date_rank", 3) in (0, 1)
    return {"nyc": 0 if preferred else 2, "remote": 1 if preferred else 4,
            "other": 3 if preferred else 5}[r["tier"]]


def order_candidates(records):
    """Priority by priority (see PRIORITIES); inside each, the four areas take
    turns, stated dates before unstated ones, newest first."""
    ordered = []
    for level in range(len(PRIORITIES)):
        queues = []
        for glabel in GROUPS.values():
            q = [r for r in records if priority(r) == level and r["group"] == glabel]
            q.sort(key=lambda r: r["posted"], reverse=True)
            q.sort(key=lambda r: r["date_rank"])  # stable
            queues.append(q)
        while any(queues):
            for q in queues:
                if q:
                    ordered.append(q.pop(0))
    return ordered


def pick(records, n):
    """n roles, one per company; a company repeats only when there are not enough others."""
    ordered = order_candidates(records)
    chosen, companies = [], set()
    for r in ordered:
        if len(chosen) == n:
            break
        c = r["company"].lower()
        if c not in companies:
            chosen.append(r)
            companies.add(c)
    for r in ordered:
        if len(chosen) == n:
            break
        if r not in chosen:
            chosen.append(r)
    # Listed priority by priority, the areas mixed (quant, data science, consulting,
    # backend, quant, ...) rather than in one block per area.
    mixed = []
    for level in range(len(PRIORITIES)):
        queues = [[r for r in chosen if priority(r) == level and r["group"] == g] for g in GROUPS.values()]
        while any(queues):
            for q in queues:
                if q:
                    mixed.append(q.pop(0))
    return mixed


# ----------------------------------------------------------------- state

def load_json(path, default):
    if not os.path.exists(path):
        return default
    with open(path, encoding="utf-8") as fh:
        return json.load(fh)


def save_json(path, obj):
    with open(path, "w", encoding="utf-8") as fh:
        json.dump(obj, fh, indent=1, sort_keys=True)


def date_tag(dt):
    """2026-9-8 style: no zero padding, dashes because slashes can't be in file names."""
    return f"{dt.year}-{dt.month}-{dt.day}"


def delete_old_reports(max_age_days):
    """Remove dated files in shortlists/ older than the cutoff. Returns names removed."""
    if not os.path.isdir(OUT_DIR):
        return []
    cutoff = datetime.now().date() - timedelta(days=max_age_days)
    removed = []
    for name in os.listdir(OUT_DIR):
        m = RE_DATED.match(name)
        if not m:
            continue
        try:
            file_date = datetime(int(m[1]), int(m[2]), int(m[3])).date()
        except ValueError:
            continue
        if file_date < cutoff:
            try:
                os.remove(os.path.join(OUT_DIR, name))
                removed.append(name)
            except OSError:
                pass  # open in Excel; it goes next time
    return removed


# ----------------------------------------------------------------- Excel output

COLUMNS = [  # header, width
    ("#", 5), ("Company", 28), ("Role", 52), ("Apply link", 70), ("Location", 34),
    ("Area", 24), ("Group", 30), ("Summer dates", 30), ("Priority", 42), ("Posted", 11), ("Pay", 18), ("Note", 40),
]
TIER_FILL = {"nyc": "FFF2CC", "remote": "DDEBF7", "other": "F2F2F2"}


def write_xlsx(path, records, meta):
    wb = Workbook()
    ws = wb.active
    ws.title = meta["tag"]  # sheet tab shows the date, same as the file name
    border = Border(bottom=Side(style="thin", color="D9D9D9"))

    # Row 1: summary line. Row 2: blank. Row 3: header. Data from row 4.
    c = meta["counts"]
    known = sum(1 for rec in records if rec.get("date_rank") in (0, 1))
    ws["A1"] = (f"{len(records)} applications for {meta['tag']}  |  {known} with stated dates inside "
                f"June 1 - Aug 31  |  {c['nyc']} Manhattan / Jersey City, "
                f"{c['remote']} remote, {c['other']} other US  |  {meta['waiting']} more eligible roles "
                f"waiting for later days  |  data fetched {meta['stamp']} UTC")
    ws["A1"].font = Font(bold=True, size=12)
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(COLUMNS))

    header_row = 3
    for col, (name, width) in enumerate(COLUMNS, start=1):
        cell = ws.cell(row=header_row, column=col, value=name)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="305496")
        cell.alignment = Alignment(vertical="center")
        ws.column_dimensions[get_column_letter(col)].width = width
    ws.row_dimensions[header_row].height = 20

    r = header_row + 1
    for n, rec in enumerate(records, start=1):
        values = [n, rec["company"], rec["title"], rec["url"], rec["location"], TIERS[rec["tier"]],
                  rec["group"], rec.get("dates") or "not stated", PRIORITIES[priority(rec)], rec["posted"], rec["salary"],
                  "; ".join(filter(None, [NOTES.get(rec["sponsorship"], "")] + rec.get("notes", [])))]
        for col, v in enumerate(values, start=1):
            cell = ws.cell(row=r, column=col, value=v)
            cell.border = border
            cell.alignment = Alignment(vertical="top", wrap_text=(col in (2, 3, 4, 5)))
        ws.cell(row=r, column=2).font = Font(bold=True)
        link = ws.cell(row=r, column=4)
        link.hyperlink = rec["url"]
        link.font = Font(color="0563C1", underline="single")
        ws.cell(row=r, column=6).fill = PatternFill("solid", fgColor=TIER_FILL[rec["tier"]])
        r += 1

    last = max(r - 1, header_row)
    ws.auto_filter.ref = f"A{header_row}:{get_column_letter(len(COLUMNS))}{last}"
    ws.freeze_panes = ws.cell(row=header_row + 1, column=1)
    wb.save(path)


# ----------------------------------------------------------------- applied log

def load_log():
    """{"roles": {role key: {date, company, title, url, location, group, ids}}}"""
    return load_json(LOG, {"roles": {}})


def log_entry(rec, day):
    return {"date": day, "company": rec["company"], "title": rec["title"], "url": rec["url"],
            "location": rec["location"], "group": rec["group"], "ids": rec["ids"]}


def prune_log(log, open_ids):
    """Drop logged roles whose posting is no longer open. Returns the entries removed."""
    removed = []
    for key in list(log["roles"]):
        entry = log["roles"][key]
        if not any(i in open_ids for i in entry["ids"]):
            removed.append(log["roles"].pop(key))
    return removed


LOG_COLUMNS = [("Listed on", 12), ("Company", 28), ("Role", 52), ("Apply link", 70),
               ("Location", 34), ("Group", 30)]


def write_log_xlsx(log, removed_now):
    """Readable copy of the log. Returns False when the file is open in Excel."""
    entries = sorted(log["roles"].values(), key=lambda e: e["company"].lower())
    entries.sort(key=lambda e: e["date"], reverse=True)  # stable: newest day first
    wb = Workbook()
    ws = wb.active
    ws.title = "Applied log"
    border = Border(bottom=Side(style="thin", color="D9D9D9"))
    ws["A1"] = (f"Applied log  |  {len(entries)} roles, none of them will be listed again  |  "
                f"{removed_now} removed on the last run because the posting closed")
    ws["A1"].font = Font(bold=True, size=12)
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(LOG_COLUMNS))
    for col, (name, width) in enumerate(LOG_COLUMNS, start=1):
        cell = ws.cell(row=3, column=col, value=name)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="305496")
        ws.column_dimensions[get_column_letter(col)].width = width
    r = 4
    for e in entries:
        y, m, d = e["date"].split("-")
        values = [f"{int(y)}-{int(m)}-{int(d)}", e["company"], e["title"], e["url"], e["location"], e["group"]]
        for col, v in enumerate(values, start=1):
            cell = ws.cell(row=r, column=col, value=v)
            cell.border = border
            cell.alignment = Alignment(vertical="top", wrap_text=(col in (2, 3, 4, 5)))
        ws.cell(row=r, column=2).font = Font(bold=True)
        link = ws.cell(row=r, column=4)
        link.hyperlink = e["url"]
        link.font = Font(color="0563C1", underline="single")
        r += 1
    ws.auto_filter.ref = f"A3:{get_column_letter(len(LOG_COLUMNS))}{max(r - 1, 3)}"
    ws.freeze_panes = "A4"
    os.makedirs(OUT_DIR, exist_ok=True)
    try:
        wb.save(LOG_XLSX)
    except PermissionError:
        print("  (applied log.xlsx is open in Excel; it will be refreshed on the next run)")
        return False
    return True


# ----------------------------------------------------------------- main

def build_pool(all_jobs, log, today, cache):
    """Eligible, in-scope summer roles that are not in the applied log, plus drop counts."""
    dropped = Counter()
    kept = []
    for j in all_jobs:
        g = classify(j)
        if not g:
            continue
        reason = why_ineligible(j, today)
        if reason:
            dropped[reason] += 1
            continue
        if (cache.get(j["id"]) or {}).get("status") == "gone":
            dropped["posting taken down (the job board no longer has it)"] += 1
            continue
        kept.append(dict(j, group_key=g))

    logged_ids = {i for e in log["roles"].values() for i in e["ids"]}
    records = []
    for r in collapse(kept):
        if r["key"] in log["roles"] or any(i in logged_ids for i in r["ids"]):
            dropped["already in the applied log"] += 1
            continue
        ok, dates, rank = term_verdict(r, cache)
        if not ok:
            dropped[dates.split(" (")[0]] += 1
            continue
        ok, notes = requirement_verdict(r, cache)
        if not ok:
            dropped[notes] += 1
            continue
        records.append({
            "dates": dates, "date_rank": rank, "notes": notes,
            "tier": tier_of(r), "group": GROUPS[r["group_key"]], "company": r["company"],
            "title": r["title"], "location": "; ".join(r["locations"]), "season": r.get("season") or "",
            "sponsorship": r.get("sponsorship") or "", "salary": r.get("salary") or "",
            "skills": ", ".join(r.get("skills") or []), "posted": posted_date(r),
            "remote": bool(r.get("remote")), "url": r["url"], "ids": r["ids"], "key": r["key"],
        })
    return records, dropped


def skip_numbers(today_state, log, numbers):
    """Take the given row numbers of today's list out of the applied log."""
    records = today_state.get("records", [])
    done = []
    for n in numbers:
        if 1 <= n <= len(records) and log["roles"].pop(records[n - 1]["key"], None):
            done.append(f"#{n} {records[n - 1]['company']}")
    save_json(LOG, log)
    write_log_xlsx(log, 0)
    print("Taken out of the applied log, so they can be listed again: " + (", ".join(done) or "nothing"))
    return 0


def main():
    dry = "--dry-run" in sys.argv
    force_new = "--new" in sys.argv
    now = datetime.now()  # local time, matches the 10 AM run
    tag = date_tag(now)
    stamp_day = now.strftime("%Y-%m-%d")

    today_state = load_json(TODAY, {})
    same_day = today_state.get("date") == tag
    if "--urls" in sys.argv:  # used by the open-in-Chrome command
        if not same_day:
            print(f"No list built for {tag} yet (latest is {today_state.get('date', 'none')}).")
            return 1
        print(f"{today_state['file']}  |  data fetched {today_state['meta']['stamp']} UTC")
        for n, r in enumerate(today_state["records"], start=1):
            print(f"{n}\t{r['company']}\t{r['url']}")
        return 0

    log = load_log()
    if "--skip" in sys.argv:
        numbers = [int(a) for a in sys.argv[sys.argv.index("--skip") + 1:] if a.isdigit()]
        return skip_numbers(today_state, log, numbers)

    d = load_json(JOBS, None)
    if d is None:
        print("No docs/api/jobs.json yet; run `run.py update` first.")
        return 1

    # Logged roles whose posting closed leave the log.
    removed_from_log = prune_log(log, {j["id"] for j in d["jobs"]})
    if removed_from_log and not dry:
        save_json(LOG, log)
    if removed_from_log:
        print(f"Applied log: {len(removed_from_log)} role(s) removed because the posting closed:")
        for e in removed_from_log:
            print(f"  - {e['company']}: {e['title'][:70]}")

    if same_day and not force_new and not dry:
        path = os.path.join(ROOT, today_state["file"])
        if not os.path.exists(path):
            write_xlsx(path, today_state["records"], today_state["meta"])
        write_log_xlsx(log, len(removed_from_log))
        print(f"Today's list is already built: {today_state['file']} ({len(today_state['records'])} roles).")
        print("Run with --new for another batch of 20.")
        return 0

    records, dropped = build_pool(d["jobs"], log, now.date(), load_json(TERM_CACHE, {}))
    chosen = pick(records, PER_DAY)
    counts = {t: sum(1 for r in chosen if r["tier"] == t) for t in TIERS}
    waiting = len(records) - len(chosen)

    print(f"{len(d['jobs'])} open postings; {len(records)} eligible roles not in the applied log.")
    for reason, n in sorted(dropped.items()):
        print(f"  dropped {n:4d}  {reason}")
    if dry:
        for n, r in enumerate(chosen, start=1):
            print(f"  {n:2d}. [{r['tier']:6s}] {r['company'][:22]:22s} {r['title'][:46]:46s} "
                  f"{r['dates'] or 'dates not stated'}")
        print("Dry run: nothing written.")
        return 0
    if not chosen:
        write_log_xlsx(log, len(removed_from_log))
        print("Nothing left to hand out today.")
        return 0

    batch = today_state.get("batch", 1) + 1 if same_day else 1
    name = f"{tag}.xlsx" if batch == 1 else f"{tag} batch {batch}.xlsx"
    rel = os.path.join("shortlists", name)
    meta = {"tag": tag, "stamp": d["generated_at"], "counts": counts, "waiting": waiting}
    os.makedirs(OUT_DIR, exist_ok=True)
    try:
        write_xlsx(os.path.join(ROOT, rel), chosen, meta)
    except PermissionError:
        print(f"Could not write {rel}: close it in Excel and run again. Nothing was handed out.")
        return 1

    for r in chosen:
        log["roles"][r["key"]] = log_entry(r, stamp_day)
    save_json(LOG, log)
    write_log_xlsx(log, len(removed_from_log))
    save_json(TODAY, {"date": tag, "batch": batch, "file": rel, "meta": meta, "records": chosen})
    save_json(LATEST, {"generated_at": d["generated_at"], "records": chosen})
    save_json(NEW, {"generated_at": d["generated_at"], "records": chosen})

    removed = delete_old_reports(KEEP_FILES_DAYS)
    print(f"Wrote {rel}: {len(chosen)} roles ({counts['nyc']} Manhattan / Jersey City, "
          f"{counts['remote']} remote, {counts['other']} other US); {waiting} more waiting."
          + (f" Deleted {len(removed)} old file(s)." if removed else ""))
    print(f"Applied log now holds {len(log['roles'])} roles.")
    if len(chosen) < PER_DAY:
        print(f"  Only {len(chosen)} eligible roles were left, fewer than {PER_DAY}.")
    return 0


if __name__ == "__main__":
    sys.exit(main())
